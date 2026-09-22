"""
activity_graph.py — agrégation et rendu d'un graphe d'activité horaire.

Source de données : xlsx d'observations Vigie-Chiro (téléchargé via
``download_observations_as_xlsx`` ou présent à la racine de la session).
On utilise en priorité le ``validateur_taxon`` (validation humaine via le logiciel externe), avec fallback sur ``observateur_taxon`` puis ``tadarida_taxon``.

L'heure de chaque contact est extraite du **nom du fichier WAV**, pas du
champ ``temps_debut`` (qui est en secondes depuis le début de l'enregistrement
et nécessiterait de recouper avec l'heure de début du WAV).

Format attendu du nom de fichier :
``Car<SITE>-<YYYY>-Pass<N>-<POINT>-<SERIAL>_<YYYYMMDD>_<HHMMSS>[_…].wav``

Ce module est autonome (pas de dépendance Tk/CTk) → testable et réutilisable.
Le rendu visuel se fait dans ``gui_activity.py``.
"""

from __future__ import annotations

import os
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime, time, timedelta
from pathlib import Path
from typing import Iterable


# ---------------------------------------------------------------------------
# Recherche dans les listes de filtres (logique pure → testable)
# ---------------------------------------------------------------------------

def _norm_search(s) -> str:
    """Normalise pour la recherche : minuscules, sans accents."""
    import unicodedata
    txt = unicodedata.normalize("NFD", str(s or ""))
    return "".join(c for c in txt if unicodedata.category(c) != "Mn").casefold()


def filter_items(items, query, *, label_fn=str) -> list:
    """Sous-ensemble d'``items`` dont le libellé contient ``query``.

    Insensible à la casse **et aux accents**. Requête vide/blanche → tout.
    Utilisé par les barres de recherche du panneau de filtres (nuits, taxons) :
    évite de faire défiler des dizaines d'items à l'aveugle.
    """
    q = _norm_search(query).strip()
    if not q:
        return list(items)
    return [it for it in items if q in _norm_search(label_fn(it))]


# Horodatage dans le nom (Vigie-Chiro / Wildlife / AudioMoth expandé).
# Pas d'ancre `$` : le tableur Tadarida n'a souvent pas l'extension, et un
# suffixe / espace en trop ne doit pas faire tomber le parse (sinon la
# date calendaire scinde une pose à minuit).
_RE_FILENAME_TIMESTAMP = re.compile(
    r"_(\d{8})_(\d{2})(\d{2})(\d{2})"
)
# Titley Swift/Ranger usine : 2026-08-21 20-45-29  (espace ou underscore)
_RE_TITLEY_TIMESTAMP = re.compile(
    r"(20\d{2})-(\d{2})-(\d{2})[ _](\d{2})[-:](\d{2})[-:](\d{2})"
)
# AudioMoth expandé : date en tête, sans underscore devant.
_RE_DATE_HEAD_TIMESTAMP = re.compile(
    r"(?<!\d)(\d{8})_(\d{2})(\d{2})(\d{2})"
)

# Regex contexte session : préfixe canonique Vigie-Chiro
#   Car<SITE>-<YYYY>-Pass<N>-<POINT>-<SERIAL>_<...>
_RE_FILENAME_CONTEXT = re.compile(
    r"^Car(?P<site>\d{5,6})-(?P<year>\d{4})-Pass(?P<passage>\d+)-"
    r"(?P<point>[A-Za-z]\d+)-(?P<serial>[^_]+)_"
)


def parse_filename_context(name: str) -> dict | None:
    """Extrait site / passage / point depuis un nom WAV canonique.

    Retourne None si le nom n'est pas au format Vigie-Chiro
    (WAV brut non renommé par exemple).
    """
    m = _RE_FILENAME_CONTEXT.match(name)
    if not m:
        return None
    d = m.groupdict()
    try:
        d["passage"] = int(d["passage"])
    except (TypeError, ValueError):
        return None
    d["site"] = str(d["site"]).zfill(6)
    d["point"] = d["point"].upper()
    return d


@dataclass
class HourlyBin:
    """Comptage de contacts validés sur une tranche horaire d'une nuit."""
    start_minutes_from_midnight: int  # 0..1440
    n_contacts: int = 0


@dataclass
class NightActivity:
    """Activité horaire pour 1 nuit, 1 espèce (ou groupe / "tous")."""
    label: str            # ex: "2025-07-12 · Pippip"
    night_date: str       # ISO YYYY-MM-DD
    taxon: str            # nom court ou "(tous)"
    bins: list[int]       # nb contacts par tranche, indexé sur _bin_index
    bin_minutes: int      # taille d'une tranche (5, 15, 30 minutes)


def _night_date_iso(date_s: str, minutes_from_midnight: int,
                     cutoff_hour: int = 12) -> str:
    """Retourne la date ISO de la NUIT à laquelle rattacher un contact.

    Un contact avant `cutoff_hour` (midi par défaut) est rattaché à la veille
    (nuit commencée la veille au soir). date_s au format YYYYMMDD.
    """
    from datetime import date, timedelta
    try:
        y, mo, d = int(date_s[:4]), int(date_s[4:6]), int(date_s[6:8])
        dd = date(y, mo, d)
    except (ValueError, IndexError):
        # Fallback : format brut sans rattachement
        return f"{date_s[:4]}-{date_s[4:6]}-{date_s[6:8]}"
    if minutes_from_midnight < cutoff_hour * 60:
        dd = dd - timedelta(days=1)
    return dd.isoformat()


def parse_filename_time(name: str) -> tuple[str, int] | None:
    """Extrait (date_yyyymmdd, minutes_from_midnight) depuis un nom de fichier.

    Retourne None si format non reconnu.
    """
    name = str(name or "").strip()
    if not name:
        return None
    m = _RE_FILENAME_TIMESTAMP.search(name)
    if m:
        date_s = m.group(1)
        h, mn, s = int(m.group(2)), int(m.group(3)), int(m.group(4))
    else:
        tm = _RE_TITLEY_TIMESTAMP.search(name)
        if tm:
            date_s = f"{tm.group(1)}{tm.group(2)}{tm.group(3)}"
            h, mn, s = int(tm.group(4)), int(tm.group(5)), int(tm.group(6))
        else:
            hm = _RE_DATE_HEAD_TIMESTAMP.search(name)
            if not hm:
                return None
            date_s = hm.group(1)
            h, mn, s = int(hm.group(2)), int(hm.group(3)), int(hm.group(4))
    if not (0 <= h <= 23 and 0 <= mn <= 59 and 0 <= s <= 59):
        return None
    return date_s, h * 60 + mn


def _header_index(header_row) -> dict[str, int]:
    """Indices de colonnes, insensible à la casse / espaces."""
    idx: dict[str, int] = {}
    for i, h in enumerate(header_row or ()):
        if h is None:
            continue
        key = str(h).strip().lower()
        if key:
            idx[key] = i
    return idx


def is_chiro_taxon(taxon: str | None) -> bool:
    """True si le code est un chiroptère (liste SpeciesList / préfixes)."""
    from taxons import classify_taxon
    return classify_taxon(taxon) == "chiros"


def _cell(row, i: int | None):
    if i is None or i < 0 or i >= len(row):
        return None
    return row[i]


def best_taxon(row_lookup: dict, idx_validateur: int | None,
                idx_observateur: int | None,
                idx_tadarida: int | None) -> str | None:
    """Retourne le taxon le plus "trustworthy" disponible pour un contact.

    Priorité : validateur (humain) > observateur (humain) > tadarida (auto).
    Retourne None si la valeur est vide ou ne peut être lue.
    """
    for idx in (idx_validateur, idx_observateur, idx_tadarida):
        if idx is None:
            continue
        v = row_lookup.get(idx)
        if v is None:
            continue
        s = str(v).strip()
        if s:
            return s
    return None


def _bin_index(minutes_from_midnight: int, bin_size: int) -> int:
    """Index de la tranche horaire (toutes alignées sur minuit, en cycles 24h)."""
    minutes_from_midnight %= 1440
    return minutes_from_midnight // bin_size


def _bin_count(bin_size: int) -> int:
    return 1440 // bin_size


def _contact_key_from_filename(fname_s: str, taxon: str, bin_minutes: int):
    parsed_time = parse_filename_time(fname_s)
    if not parsed_time:
        return None
    date_s, mins = parsed_time
    night_date = _night_date_iso(date_s, mins)
    ctx = parse_filename_context(fname_s)
    if ctx is not None:
        site, point, passage = ctx["site"], ctx["point"], ctx["passage"]
    else:
        site, point, passage = "?", "?", None
    key = (site, point, passage, night_date, taxon)
    return key, mins


def aggregate_rows(headers, rows, *,
                   bin_minutes: int = 30,
                   taxon_filter: str | None = None,
                   use_only_validated: bool = False,
                   use_observer_taxon: bool = False,
                   chiros_only: bool = False,
                   use_mnhn: bool = False,
                   ) -> dict[tuple[str, str, int | None, str, str], list[int]]:
    """Agrège des lignes d'observations en
    ``{(site, point, passage, night_date, taxon): bins}``.

    - ``use_only_validated`` : ignore les contacts sans identification
      humaine (``validateur_taxon`` **ou** ``observateur_taxon``).
    - ``use_observer_taxon`` : ne garder que les lignes avec
      ``observateur_taxon`` et grouper sous ce code (issue #4.10).
    - ``chiros_only`` : ignorer orthoptères / bruit / oiseaux (issue #4.9).
    - ``use_mnhn`` : relecture d’un ``_Vu`` produit dans le logiciel externe
      (même règle que la Synthèse). Ce calcul n’est pas l’évaluation
      d’activité du logiciel externe.
      Prime sur ``use_only_validated`` et ``use_observer_taxon``.
    """
    if bin_minutes <= 0 or 1440 % bin_minutes != 0:
        raise ValueError("bin_minutes doit diviser 1440 (1, 5, 15, 30, 60…)")
    idx = _header_index(headers)
    idx_filename = idx.get("nom du fichier")
    if idx_filename is None:
        idx_filename = idx.get("fichier")
    if idx_filename is None:
        idx_filename = idx.get("filename")
    idx_validateur = idx.get("validateur_taxon")
    idx_observateur = idx.get("observateur_taxon")
    idx_tadarida = idx.get("tadarida_taxon")
    if idx_filename is None:
        return {}

    nbins = _bin_count(bin_minutes)
    result: dict[tuple[str, str, int | None, str, str], list[int]] = {}

    def _add(fname_s: str, taxon: str):
        if taxon_filter and taxon != taxon_filter:
            return
        if chiros_only and not is_chiro_taxon(taxon):
            return
        parsed = _contact_key_from_filename(fname_s, taxon, bin_minutes)
        if parsed is None:
            return
        key, mins = parsed
        if key not in result:
            result[key] = [0] * nbins
        result[key][_bin_index(mins, bin_minutes)] += 1

    if use_mnhn:
        from synthesis import iter_mnhn_contacts
        for r, taxon in iter_mnhn_contacts(headers, rows):
            if not r or len(r) <= idx_filename:
                continue
            fname = r[idx_filename]
            if not fname:
                continue
            _add(str(fname), taxon)
        return result

    for r in rows:
        if not r or len(r) <= idx_filename:
            continue
        fname = r[idx_filename]
        if not fname:
            continue
        fname_s = str(fname)

        obs = _cell(r, idx_observateur)
        obs_s = str(obs).strip() if obs not in (None, "") else ""
        val = _cell(r, idx_validateur)
        val_s = str(val).strip() if val not in (None, "") else ""

        if use_observer_taxon:
            if not obs_s:
                continue
            taxon = obs_s
        else:
            if use_only_validated and not (val_s or obs_s):
                continue
            row_lookup = {i: r[i] if i < len(r) else None for i in range(len(r))}
            taxon = best_taxon(row_lookup, idx_validateur,
                                idx_observateur, idx_tadarida) or "(?)"
        _add(fname_s, taxon)

    return result


def is_observations_xlsx_name(name: str) -> bool:
    """True pour un tableur Vigie-Chiro d'observations (pas un backup)."""
    low = str(name or "").lower()
    if not low.endswith(".xlsx"):
        return False
    if "_cleanup" in low or "_backup" in low:
        return False
    return "observations" in low and "participation-" in low


def is_vu_csv_name(name: str) -> bool:
    """True pour un fichier annexe ``*_Vu.csv``."""
    low = str(name or "").lower()
    if "_cleanup" in low or "_backup" in low:
        return False
    return low.endswith("_vu.csv")


_SKIP_DIR_NAMES = {".git", "__pycache__", "build", "dist"}
# Dossiers d'audio : on lit les fichiers du niveau (xlsx / _Vu), on ne
# descend pas dans des milliers de WAV (issue #10, lenteur du 1er scan).
_LEAF_DIR_NAMES = {"data", "data_k"}


def discover_activity_sources(workspace: Path, *, max_depth: int = 6
                              ) -> list[Path]:
    """xlsx ``participation-*-observations`` + CSV ``_Vu`` sous le workspace.

    Un ``_Vu`` est accepté où le logiciel externe / l'utilisateur le pose (chirosurf/,
    Data_k/, Data/, racine de session). Pas seulement dans trois dossiers
    magiques : sinon un ``_Vu`` collé à côté de l'xlsx disparaît du graphe
    alors que la Synthèse (par session) le voit encore.
    """
    workspace = Path(workspace)
    if not workspace.is_dir():
        return []
    xlsx: list[Path] = []
    vu_csvs: list[Path] = []
    try:
        for dirpath, dirnames, filenames in os.walk(workspace):
            try:
                rel = Path(dirpath).relative_to(workspace)
                parts = () if rel.as_posix() == "." else rel.parts
            except ValueError:
                dirnames[:] = []
                continue
            depth = len(parts)
            base = Path(dirpath).name.lower()
            keep: list[str] = []
            if depth < max_depth and base not in _LEAF_DIR_NAMES:
                for d in dirnames:
                    if d.startswith(".") or d.lower() in _SKIP_DIR_NAMES:
                        continue
                    keep.append(d)
            dirnames[:] = keep
            for name in filenames:
                p = Path(dirpath) / name
                if is_vu_csv_name(name):
                    vu_csvs.append(p)
                elif is_observations_xlsx_name(name):
                    xlsx.append(p)
    except OSError:
        return []
    vu_csvs.sort(key=lambda q: (
        0 if q.parent.name.lower() in ("chirosurf", "chirosurf_nuits") else 1,
        str(q).lower(),
    ))
    vu_kept: list[Path] = []
    seen_vu: set[tuple] = set()
    try:
        from chirosurf_nights import parse_chirosurf_csv_name
    except Exception:
        parse_chirosurf_csv_name = lambda _n: None  # noqa: E731
    for vp in vu_csvs:
        parsed = parse_chirosurf_csv_name(vp.name)
        idx = parsed[0] if parsed else None
        parent = vp.parent
        session = parent.parent if parent.name.lower() in (
            "chirosurf", "chirosurf_nuits", "data_k", "data",
        ) else parent
        key = (str(session).lower(), idx if idx is not None else vp.name.lower())
        if key in seen_vu:
            continue
        seen_vu.add(key)
        vu_kept.append(vp)
    return sorted(xlsx + vu_kept, key=lambda q: str(q).lower())


def load_observation_table(path: Path) -> tuple[list, list] | None:
    """Charge un xlsx / CSV en listes (headers, rows). None si illisible.

    Materialise les lignes : un iterateur openpyxl meurt à la fermeture du
    classeur, et on veut relire depuis la RAM à chaque filtre.
    """
    path = Path(path)
    try:
        if not path.is_file():
            return None
    except OSError:
        return None
    suffix = path.suffix.lower()
    if suffix == ".csv":
        import csv
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError:
            return None
        lines = text.splitlines()
        if not lines:
            return None
        first = lines[0]
        delim = ";" if first.count(";") >= first.count(",") else ","
        reader = csv.reader(lines, delimiter=delim)
        headers = next(reader, None)
        if not headers:
            return None
        rows = [list(r) for r in reader if r and any(str(c).strip() for c in r)]
        return list(headers), rows
    if suffix != ".xlsx":
        return None
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return None
    try:
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            return None
        headers = list(header)
        rows = []
        for r in it:
            if r and any(c is not None and str(c).strip() for c in r):
                rows.append(list(r))
        return headers, rows
    except Exception:
        return None
    finally:
        try:
            wb.close()
        except Exception:
            pass


class ObservationTableCache:
    """Cache (mtime, taille) → (headers, rows) pour ne plus relire le disque."""

    def __init__(self):
        self._data: dict[str, tuple[float, int, list, list]] = {}

    def load(self, path: Path) -> tuple[list, list] | None:
        path = Path(path)
        try:
            st = path.stat()
        except OSError:
            return None
        key = str(path)
        cached = self._data.get(key)
        if cached and cached[0] == st.st_mtime and cached[1] == st.st_size:
            return cached[2], cached[3]
        loaded = load_observation_table(path)
        if loaded is None:
            self._data.pop(key, None)
            return None
        headers, rows = loaded
        self._data[key] = (st.st_mtime, st.st_size, headers, rows)
        return headers, rows

    def drop_missing(self, keep: Iterable[Path]) -> None:
        keep_keys = {str(Path(p)) for p in keep}
        for key in list(self._data):
            if key not in keep_keys:
                del self._data[key]


def _iter_table_file(path: Path):
    """Yield (headers, rows) depuis un xlsx ou un CSV par nuit."""
    loaded = load_observation_table(Path(path))
    if loaded is None:
        return
    yield loaded


def aggregate_xlsx(xlsx_path: Path, *,
                    bin_minutes: int = 30,
                    taxon_filter: str | None = None,
                    use_only_validated: bool = False,
                    use_observer_taxon: bool = False,
                    chiros_only: bool = False,
                    use_mnhn: bool = False,
                    ) -> dict[tuple[str, str, int | None, str, str], list[int]]:
    """Agrège un xlsx **ou** un CSV ``_Vu`` par nuit."""
    result: dict[tuple[str, str, int | None, str, str], list[int]] = {}
    for headers, rows in _iter_table_file(Path(xlsx_path)):
        partial = aggregate_rows(
            headers, rows,
            bin_minutes=bin_minutes,
            taxon_filter=taxon_filter,
            use_only_validated=use_only_validated,
            use_observer_taxon=use_observer_taxon,
            chiros_only=chiros_only,
            use_mnhn=use_mnhn,
        )
        for k, bins in partial.items():
            if k not in result:
                result[k] = list(bins)
            else:
                cur = result[k]
                for i, n in enumerate(bins):
                    if i < len(cur):
                        cur[i] += n
                    else:
                        cur.append(n)
    return result


def _activity_cover_key(key: tuple) -> tuple:
    """Clé de nuit (sans taxon) pour dédupliquer ``_Vu`` vs xlsx."""
    if len(key) >= 5:
        return key[:4]
    return key[:-1] if len(key) >= 2 else key


def _add_activity_partial(out: dict, key: tuple, bins: list[int]) -> None:
    if key not in out:
        out[key] = list(bins)
        return
    cur = out[key]
    for i, n in enumerate(bins):
        if i < len(cur):
            cur[i] += n
        else:
            cur.append(n)


def aggregate_loaded_tables(loaded: Iterable[tuple[Path, list, list]],
                            **kwargs
                            ) -> dict[tuple, list[int]]:
    """Agrège des tables déjà en mémoire. Un ``_Vu`` prime sur sa nuit."""
    items = [(Path(p), headers, rows) for p, headers, rows in loaded]
    csvs = [(p, h, r) for p, h, r in items if p.suffix.lower() == ".csv"]
    others = [(p, h, r) for p, h, r in items if p.suffix.lower() != ".csv"]
    out: dict[tuple, list[int]] = {}
    covered: set[tuple] = set()

    for _p, headers, rows in csvs:
        try:
            # La couverture appartient à la source, même si le filtre MNHN
            # ne retient aucun contact de cette nuit.
            coverage = aggregate_rows(headers, rows)
            partial = aggregate_rows(headers, rows, **kwargs)
        except Exception:
            continue
        covered.update(_activity_cover_key(k) for k in coverage)
        for k, bins in partial.items():
            _add_activity_partial(out, k, bins)
            covered.add(_activity_cover_key(k))

    for _p, headers, rows in others:
        try:
            partial = aggregate_rows(headers, rows, **kwargs)
        except Exception:
            continue
        for k, bins in partial.items():
            if _activity_cover_key(k) in covered:
                continue
            _add_activity_partial(out, k, bins)
    return out


def aggregate_multi_xlsx(paths: Iterable[Path], **kwargs
                          ) -> dict[tuple[str, str], list[int]]:
    """Agrège plusieurs xlsx / CSV ``_Vu`` en un seul dict cumulé.

    Un ``_Vu`` prime sur l'xlsx pour la même nuit (même site / point /
    passage / date). Les autres nuits de l'xlsx restent. Évite de perdre
    une nuit 2 quand seul un ``_Vu`` nuit 1 existe.
    """
    loaded: list[tuple[Path, list, list]] = []
    for p in paths:
        pair = load_observation_table(Path(p))
        if pair is None:
            continue
        loaded.append((Path(p), pair[0], pair[1]))
    return aggregate_loaded_tables(loaded, **kwargs)


def list_taxons(aggregated: dict, *, min_total: int = 1
                  ) -> list[tuple[str, int]]:
    """Liste des taxons avec nb total de contacts, triée décroissant."""
    totals: dict[str, int] = defaultdict(int)
    for key, bins in aggregated.items():
        # Support les anciennes clés (night, taxon) et les nouvelles
        # (site, point, passage, night, taxon)
        taxon = key[-1]
        totals[taxon] += sum(bins)
    return sorted(
        ((t, n) for t, n in totals.items() if n >= min_total),
        key=lambda x: x[1], reverse=True,
    )


def list_nights(aggregated: dict) -> list[str]:
    """Liste des nuits présentes (dates ISO uniques), triée."""
    # La nuit est en avant-dernière position (avant le taxon)
    nights = set()
    for key in aggregated.keys():
        nights.add(key[-2])
    return sorted(nights)


def list_sites(aggregated: dict) -> list[str]:
    """Liste des sites (carrés STOC) présents, triée.

    Inclut "?" si des fichiers non-canoniques sont présents → l'utilisateur
    voit qu'il y a des contacts non rattachés à un site précis et peut les
    cocher s'il veut quand même les inclure.
    """
    sites = set()
    for key in aggregated.keys():
        if len(key) >= 5:
            site = key[0]
            if site:
                sites.add(site)
    # "?" en dernier pour ne pas polluer le haut de la liste
    real = sorted(s for s in sites if s != "?")
    if "?" in sites:
        real.append("?")
    return real


def list_points(aggregated: dict, site_filter: str | None = None
                  ) -> list[str]:
    """Liste des points présents (filtrable par site)."""
    points = set()
    for key in aggregated.keys():
        if len(key) < 5:
            continue
        site, point = key[0], key[1]
        if site_filter and site != site_filter:
            continue
        if point:
            points.add(point)
    real = sorted(p for p in points if p != "?")
    if "?" in points:
        real.append("?")
    return real


def list_passages(aggregated: dict) -> list[int | None]:
    """Liste des passages (1, 2, …) présents. Inclut None si des contacts
    n'ont pas de passage extractible (fichiers non canoniques)."""
    passages = set()
    for key in aggregated.keys():
        if len(key) < 5:
            continue
        passages.add(key[2])
    real = sorted(p for p in passages if p is not None)
    if None in passages:
        real.append(None)  # type: ignore[arg-type]
    return real


def _sorted_with_placeholders(items) -> list:
    """Trie en repoussant les placeholders ("?" / None) en fin de liste."""
    real = sorted(x for x in items if x not in (None, "?"))
    if "?" in items:
        real.append("?")
    if None in items:
        real.append(None)
    return real


def cascade_options(aggregated: dict, *, sel_sites=None, sel_points=None,
                    sel_passages=None) -> dict:
    """Options disponibles pour chaque dimension, EN CASCADE.

    Chaque dimension aval n'inclut que les valeurs compatibles avec les
    sélections amont fournies (``None`` = pas de contrainte sur cette dimension) :
      - ``sites``    : tous les sites présents ;
      - ``points``   : uniquement dans les sites sélectionnés ;
      - ``passages`` : uniquement dans (sites + points) sélectionnés ;
      - ``nights``   : uniquement dans (sites + points + passages) sélectionnés.

    Retourne ``{"sites", "points", "passages", "nights"}`` (listes triées,
    placeholders "?"/None en fin). Fonction PURE → testable.
    """
    sites, points, passages, nights = set(), set(), set(), set()
    for key in aggregated.keys():
        if len(key) < 5:
            continue
        s, p, pa, ni, _tx = key
        if s:
            sites.add(s)
        if sel_sites is not None and s not in sel_sites:
            continue
        if p:
            points.add(p)
        if sel_points is not None and p not in sel_points:
            continue
        passages.add(pa)
        if sel_passages is not None and pa not in sel_passages:
            continue
        if ni:
            nights.add(ni)
    return {
        "sites": _sorted_with_placeholders(sites),
        "points": _sorted_with_placeholders(points),
        "passages": _sorted_with_placeholders(passages),
        "nights": sorted(nights),
    }
